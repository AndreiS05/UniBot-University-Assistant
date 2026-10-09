# -*- coding: utf-8 -*-
"""
Construieste baza de date facultate.db din fisierul UniBot_Date_UVT.xlsx.
Ruleaza:  python setup_db.py
"""
import sqlite3
import os
import shutil
import tempfile
import unicodedata
from openpyxl import load_workbook

XLSX = "UniBot_Date_UVT.xlsx"


def open_workbook(path):
    """Deschide Excel-ul; daca e blocat (deschis in Excel), citeste dintr-o copie."""
    try:
        return load_workbook(path, data_only=True)
    except PermissionError:
        tmp = os.path.join(tempfile.gettempdir(), "unibot_date_tmp.xlsx")
        shutil.copyfile(path, tmp)
        print("(Excel deschis - citesc dintr-o copie temporara)")
        return load_workbook(tmp, data_only=True)


def val(x):
    """Curata o celula: None/'' -> None, altfel valoarea."""
    if x is None:
        return None
    s = str(x).strip()
    return s if s else None


def to_int(x):
    v = val(x)
    if v is None:
        return None
    try:
        return int(float(v))
    except ValueError:
        return None


INFORMATII_GENERALE = [
    ("absolvire",
     "absolv, termin facultat, promov facultat, conditii absolvire, cum absolv, ce trebuie sa termin facultatea, cum se termina facultatea",
     "Ca să absolvi facultatea (specializarea Informatică) trebuie să îndeplinești toate condițiile de mai jos:\n"
     "• promovezi toate materiile obligatorii și materiile opționale alese;\n"
     "• acumulezi minimul de credite ECTS: minim 60 pe fiecare an și 30 pe fiecare semestru (180 în total);\n"
     "• pentru a trece din anul 2 în anul 3 ai nevoie de minim 100 de credite dacă ai promovat toate materiile din anul 1, sau de minim 110 credite dacă ai rămas cu vreo materie nepromovată din anul 1;\n"
     "• realizezi 4 din cele 6 semestre de educație fizică și 4 de limbă străină (engleza fiind obligatorie dacă nu deții certificat Cambridge);\n"
     "• completezi 3 discipline complementare (DCT), de la alte facultăți din UVT, începând din anul 2;\n"
     "• promovezi examenul final de licență."),
    ("promovare_an",
     "trec anul, trecere anul, promovare anul, cate credite anul, sa trec in anul, promovez anul, din anul 2 in anul 3, cate credite imi trebuie sa trec",
     "Pentru a promova din anul 2 în anul 3 ai nevoie de minim 100 de credite dacă ai promovat toate materiile din anul 1, sau de minim 110 credite dacă ai rămas cu vreo materie nepromovată din anul 1."),
    ("examen_licenta",
     "examen licent, de licenta, grile, examen final, cum se da licenta, cum promovez licenta, proba scrisa licenta",
     "Examenul final de licență are două probe: (1) o probă scrisă cu 30 de întrebări grilă din materiile primilor 2 ani de studiu și (2) prezentarea proiectului final (lucrarea de licență). Fiecare probă trebuie să aibă nota minimă 5, iar nota finală de licență (rezultată din cele două) trebuie să fie cel puțin 6 pentru ca examenul de licență să fie promovat."),
    ("dct",
     "dct, discipline complementar, disciplinele complementar, complementar alte facultat, ce sunt disciplinele complementare",
     "Disciplinele complementare (DCT) sunt 3 discipline pe care trebuie să le completezi de-a lungul celor 3 ani, începând cu anul 2. Sunt materii de la alte facultăți din cadrul UVT, menite să diversifice modul de învățare al studenților. Pentru mai multe detalii despre aceste discipline, întreabă la infocentru sau consultă site-ul facultății."),
    # Admitere (sursa: info.uvt.ro/admitere-licenta; date orientative pentru 2026, de verificat anual)
    ("admitere_acte",
     "acte admitere, acte necesare, documente admitere, ce acte, dosar admitere, ce documente imi trebuie",
     "Actele necesare pentru înscrierea la admitere sunt:\n"
     "• carte de identitate / buletin;\n• certificat de naștere;\n• certificat de căsătorie (dacă este cazul);\n"
     "• diploma de bacalaureat sau o diplomă echivalentă;\n• adeverință medicală (nu mai veche de 90 de zile);\n"
     "• adeverință pentru alte studii universitare parcurse (dacă este cazul);\n• documente de echivalare, pentru studii efectuate în străinătate.\n"
     "Dacă vrei să afli mai multe detalii, accesează https://info.uvt.ro/admitere-licenta."),
    ("admitere_calendar",
     "cand admiterea, calendar admitere, cand inscriere, perioada admitere, cand inscrierile, data admitere, cand se da admiterea",
     "Calendarul admiterii (2026):\n"
     "• Sesiunea iulie: înscriere până pe 20 iulie 2026 (ora 14:00), proba de admitere pe 21 iulie, rezultate finale pe 26 iulie 2026.\n"
     "• Sesiunea septembrie: înscriere până pe 9 septembrie 2026 (ora 12:00), proba de admitere pe 10 septembrie, rezultate finale pe 13 septembrie 2026.\n"
     "Dacă vrei să afli mai multe detalii, accesează https://info.uvt.ro/admitere-licenta."),
    ("admitere",
     "admitere, cum intru, cum se intra, criteriu admitere, medie admitere, cum ma inscriu, formula admitere, cum se calculeaza media, taxa inscriere, taxa inmatriculare, cate locuri",
     "Admiterea la programul de licență Informatică (limba română) se face pe bază de concurs. "
     "Media de admitere se calculează după formula N = (max(N1, N2) + N3) / 2, unde:\n"
     "• N1 = media generală de la examenul de bacalaureat;\n"
     "• N2 = cea mai mare notă obținută la Matematică sau Informatică la bacalaureat (sau 4 dacă nu ai susținut niciuna);\n"
     "• N3 = nota la proba scrisă de admitere.\n"
     "Media de admitere trebuie să fie cel puțin 6,00, iar fiecare notă cel puțin 5,00. "
     "Programul Informatică – limba română are 200 de locuri, iar taxa de înscriere este 200 lei (înmatriculare 250 lei).\n\n"
     "Te pot ajuta și cu alte informații despre admitere: actele necesare, calendarul admiterii sau taxele de școlarizare — întreabă-mă oricând!\n\n"
     "Detalii complete găsești pe https://info.uvt.ro/admitere-licenta."),
    ("taxe",
     "taxa scolarizare, taxa de studiu, taxa pe an, cat costa facultatea, taxa anuala, taxa la facultate, taxa, taxe, taxele, ce taxe",
     "Taxa de școlarizare la Facultatea de Informatică (pentru locurile cu taxă) este:\n"
     "• anul 1: 4900 lei;\n• anul 2: 3500 lei;\n• anul 3: 3500 lei.\n"
     "Taxa se achită pe an de studiu.\n"
     "Alte taxe uzuale:\n"
     "• restanță (reexaminare): 100 lei per examen;\n"
     "• recontractarea unei discipline: 100 lei × numărul de credite al disciplinei (100 lei/credit).\n"
     "Dacă vrei să afli mai multe detalii, accesează https://info.uvt.ro."),
    # Structura anului universitar (sursa: structura oficiala UVT 2025-2026, de verificat anual)
    ("sesiuni",
     "sesiun, examene, cand e sesiunea, cand incepe sesiunea, sesiune de examene, cand sunt examenele, perioada sesiune, sesiunea de iarna, sesiunea de vara, sesiunea a, sesiunea b, sesiunea c, restante, mariri de note",
     "Sesiunile de examene în anul universitar 2025-2026 sunt:\n"
     "Semestrul I (iarnă):\n"
     "• Sesiunea A-I (principală): 17.01.2026 – 08.02.2026;\n"
     "• Sesiunea B-I (restanțe și măriri de note, gratuită): 16.02.2026 – 22.02.2026.\n"
     "Semestrul II (vară):\n"
     "• Sesiunea A-II (principală): 06.06.2026 – 28.06.2026;\n"
     "• Sesiunea B-II (restanțe și măriri de note, gratuită): 02.07.2026 – 08.07.2026.\n"
     "Sesiunile A și B au aceleași date pentru toți anii de studiu. Sesiunea C (cu taxă, restanțe din ambele semestre) se organizează doar pentru anul 3 (terminal): 22.07.2026 – 25.07.2026.\n"
     "Pentru anul 3 (terminal), examenul de licență se susține în 13.07.2026 – 16.07.2026.\n"
     "Datele se pot modifica de la an la an; pentru structura oficială, consultă site-ul UVT."),
    ("vacante",
     "vacant, cand e vacanta, vacanta de iarna, vacanta de vara, vacanta de paste, cand incepe vacanta, in ce data incepe vacanta",
     "Vacanțele din anul universitar 2025-2026 sunt:\n"
     "• Vacanța de iarnă: 20.12.2025 – 07.01.2026;\n• Vacanța de Paște: 04.04.2026 – 14.04.2026;\n• Vacanța de vară: 09.07.2026 – 04.10.2026.\n"
     "Mai există și pauze scurte între sesiuni: 09.02.2026 – 15.02.2026 (după sesiunea de iarnă) și 29.06.2026 – 01.07.2026 (între sesiunile de vară A-II și B-II).\n"
     "Datele se pot modifica de la an la an; pentru structura oficială, consultă site-ul UVT."),
    ("structura_an",
     "cate saptamani, saptamani de studiu, saptamani semestru, structura anului, structura an, structurat, cum e structurat, an universitar, organizarea anului, cum e organizat anul, cum e impartit anul, cat dureaza semestrul, cand incepe anul, incepere an, anul terminal, an terminal, cand termina anul 3, cand termina anul terminal",
     "Anul universitar 2025-2026 începe pe 29 septembrie 2025. Structura diferă între anii 1-2 și anul 3 (terminal):\n\n"
     "Anii 1 și 2 — 14 săptămâni de activitate didactică pe semestru:\n"
     "• Semestrul I: 29.09.2025 – 16.01.2026 (cu vacanța de iarnă 20.12.2025 – 07.01.2026);\n"
     "• Semestrul II: 23.02.2026 – 05.06.2026 (cu vacanța de Paște 04.04 – 14.04.2026).\n\n"
     "Anul 3 (terminal) — semestrul II are doar 12 săptămâni, programate în 4 zile pe săptămână (a 5-a zi e dedicată lucrării de licență), așa că se termină mai devreme:\n"
     "• Semestrul I: 29.09.2025 – 16.01.2026 (la fel ca la ceilalți ani);\n"
     "• Semestrul II: 23.02.2026 – 22.05.2026.\n"
     "Ceremonia de absolvire UVT: 23.05.2026; pregătirea lucrării de licență: 23.05 – 05.06.2026; examenul de licență: 13.07 – 16.07.2026.\n\n"
     "Datele se pot modifica de la an la an; pentru structura oficială, consultă site-ul UVT."),
    # Resurse pentru studenti
    ("resurse",
     "resurse, elearning, e learning, classroom, google classroom, platforma, biblioteca, unde gasesc cursurile, materiale, unde invat, unde sunt cursurile, carti, site facultate, site uvt, site oficial",
     "Resurse utile pentru studenți:\n"
     "• Platforma e-learning (https://elearning.e-uvt.ro/): aici găsești toate informațiile despre cursurile la care ai acces. Trebuie să te conectezi cu emailul tău instituțional UVT.\n"
     "• Google Classroom (https://classroom.google.com/): pentru cursuri și materiale — conectează-te cu emailul tău UVT.\n"
     "• Biblioteca UVT: se află lângă instituția principală a UVT; aici găsești cărți, spații de studiu individual și calculatoare dedicate.\n"
     "• Site-ul Facultății de Informatică: https://info.uvt.ro/\n"
     "• Site-ul principal al Universității de Vest din Timișoara: https://uvt.ro"),
]


def setup_database():
    # Verificam fisierul-sursa INAINTE de a atinge baza de date. Altfel am sterge tabelele
    # existente si am ramane cu o baza de date goala (fara materii), desi facultate.db era buna.
    if not os.path.exists(XLSX):
        print(f"EROARE: nu am gasit fisierul-sursa {XLSX}.")
        print("Baza de date NU a fost modificata.")
        print("Daca ai deja facultate.db construita, poti rula direct aplicatia: python app.py")
        return

    conn = sqlite3.connect("facultate.db")
    cursor = conn.cursor()

    cursor.execute("DROP TABLE IF EXISTS materii")

    cursor.execute("""
        CREATE TABLE materii (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            cod TEXT,
            nume TEXT NOT NULL,
            abreviere TEXT,
            an INTEGER,
            semestru INTEGER,
            categorie TEXT,
            tip TEXT,
            credite INTEGER,
            ore_curs INTEGER,
            ore_seminar INTEGER,
            ore_laborator INTEGER,
            titular TEXT,
            promovare TEXT,
            specializare TEXT DEFAULT 'Informatica'
        )
    """)

    # Tabel prezenta: cate prezente obligatorii sunt la curs / laborator per materie
    cursor.execute("DROP TABLE IF EXISTS prezenta")
    cursor.execute("""
        CREATE TABLE prezenta (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            materie_id INTEGER UNIQUE,
            prezente_curs TEXT,
            prezente_seminar TEXT,
            prezente_laborator TEXT,
            FOREIGN KEY (materie_id) REFERENCES materii(id)
        )
    """)

    # Tabel informatii: raspunsuri la intrebari generale (absolvire, licenta, DCT...)
    cursor.execute("DROP TABLE IF EXISTS informatii")
    cursor.execute("""
        CREATE TABLE informatii (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            subiect TEXT,
            cuvinte_cheie TEXT,
            raspuns TEXT
        )
    """)
    cursor.executemany(
        "INSERT INTO informatii (subiect, cuvinte_cheie, raspuns) VALUES (?,?,?)",
        INFORMATII_GENERALE,
    )

    wb = open_workbook(XLSX)
    ws = wb["Materii"]

    # header -> index
    header = [c.value for c in ws[1]]
    idx = {name: i for i, name in enumerate(header)}

    def cell(row, col_name):
        i = idx.get(col_name)
        return row[i] if i is not None and i < len(row) else None

    n = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        nume = val(cell(row, "Nume materie"))
        if not nume:
            continue
        cursor.execute("""
            INSERT INTO materii
            (cod, nume, abreviere, an, semestru, categorie, tip, credite,
             ore_curs, ore_seminar, ore_laborator, titular, promovare)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
        """, (
            val(cell(row, "Cod")), nume, val(cell(row, "Abreviere")),
            to_int(cell(row, "An")), to_int(cell(row, "Sem")),
            val(cell(row, "Categorie")), val(cell(row, "Tip")),
            to_int(cell(row, "Credite")),
            to_int(cell(row, "Ore curs")), to_int(cell(row, "Ore sem")), to_int(cell(row, "Ore lab")),
            val(cell(row, "Titular (din fise)")),
            val(cell(row, "Mod de promovare (din fise)")),
        ))
        n += 1

    # Populam prezenta cu regula generala (curs 50% = 7 din 14, seminar 10 din 14,
    # laborator 100%), aplicata doar activitatilor pe care le are disciplina, plus
    # cateva exceptii per disciplina (unde prezenta la curs nu se cere sau difera).
    def _fara_diacritice(text):
        return unicodedata.normalize("NFKD", text or "").encode("ascii", "ignore").decode().lower()

    for materie_id, nume_m, oc, os_, ol in cursor.execute(
            "SELECT id, nume, ore_curs, ore_seminar, ore_laborator FROM materii").fetchall():
        prezente_curs = "minim 7 prezențe din 14 (50%)" if oc else None
        prezente_seminar = "minim 10 prezențe din 14" if os_ else None
        prezente_laborator = "toate prezențele (100%)" if ol else None
        nn = _fara_diacritice(nume_m)
        if "sisteme de operare" in nn or nn.startswith("metode numerice"):
            prezente_curs = "nu este obligatorie"
        elif nn.startswith("inginerie"):
            prezente_curs = "minim 10 prezențe obligatorii"
        cursor.execute(
            "INSERT INTO prezenta (materie_id, prezente_curs, prezente_seminar, prezente_laborator) "
            "VALUES (?,?,?,?)",
            (materie_id, prezente_curs, prezente_seminar, prezente_laborator),
        )

    conn.commit()
    conn.close()
    print(f"Baza de date creata! {n} materii importate din {XLSX}.")


if __name__ == "__main__":
    setup_database()
